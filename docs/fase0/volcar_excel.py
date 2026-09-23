"""Vuelca los libros de Excel del despacho a texto (celda, valor o fórmula, valor en caché).

Uso, desde docs/fase0:
    python volcar_excel.py /ruta/a/Avaluos
Genera dump/<FORMATO>/<NN>_<hoja>.txt. La carpeta dump/ no se versiona: contiene datos del cliente.
"""
import glob, os, re, sys, warnings
import openpyxl

warnings.filterwarnings("ignore")
SHORT = {
    "AVALUO DPTOS ARANDAS": "REAL_ARANDAS",
    "FORMATO AVALUO PT-MEH MAQ EQ": "MEH",
    "FORMATO AVALUO PT-TCH. TERR URB CONST HAB": "TCH",
    "FORMATO AVALUO PT-TR. TERR RURAL": "TR",
    "FORMATO AVALUO PT-TRC. TERR RURAL CONST": "TRC",
    "FORMATO AVALUO PT-TU. TERR URBANO OFICIAL": "TU_OFICIAL",
    "FORMATO AVALUO PT-TU. TERR URBANO": "TU",
}

def main(src):
    for f in sorted(glob.glob(os.path.join(src, "*.xls*"))):
        base = os.path.splitext(os.path.basename(f))[0]
        tag = SHORT.get(base, re.sub(r"[^A-Za-z0-9]+", "_", base))
        kv = f.endswith("xlsm")
        wf = openpyxl.load_workbook(f, keep_vba=kv)
        wv = openpyxl.load_workbook(f, data_only=True, keep_vba=kv)
        os.makedirs(f"dump/{tag}", exist_ok=True)
        for i, ws in enumerate(wf.worksheets):
            if ws.title in ("SystemConfig", "SysAppInfo"):
                continue
            vs = wv[ws.title]
            name = re.sub(r"[^A-Za-z0-9]+", "_", ws.title).strip("_")
            with open(f"dump/{tag}/{i:02d}_{name}.txt", "w") as out:
                out.write(f"# {base} :: {ws.title} (estado={ws.sheet_state})\n")
                for row in ws.iter_rows():
                    for c in row:
                        v = c.value
                        if v is None:
                            continue
                        if isinstance(v, str) and v.startswith("="):
                            out.write(f"{c.coordinate}\t{v}\t=> {vs[c.coordinate].value!r}\n")
                        else:
                            out.write(f"{c.coordinate}\t{str(v).replace(chr(10), ' ')[:300]}\n")
        print(tag)

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Documents/Avaluos"))
