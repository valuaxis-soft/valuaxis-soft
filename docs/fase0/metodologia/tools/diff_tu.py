import openpyxl
A = '/Users/brangarciaramos/Documents/Avaluos/FORMATO AVALUO PT-TU. TERR URBANO.xlsx'
B = '/Users/brangarciaramos/Documents/Avaluos/FORMATO AVALUO PT-TU. TERR URBANO OFICIAL.xlsm'
wa = openpyxl.load_workbook(A); wb = openpyxl.load_workbook(B, keep_vba=False)
va = openpyxl.load_workbook(A, data_only=True); vb = openpyxl.load_workbook(B, data_only=True)
print('A sheets', [(s.title, s.sheet_state) for s in wa.worksheets])
print('B sheets', [(s.title, s.sheet_state) for s in wb.worksheets])
n = 0
for s in wa.sheetnames:
    if s not in wb.sheetnames: print('SOLO EN TU:', s); continue
    a, b = wa[s], wb[s]
    coords = set()
    for ws in (a, b):
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None: coords.add(c.coordinate)
    for co in sorted(coords, key=lambda x: (openpyxl.utils.cell.coordinate_from_string(x)[1], openpyxl.utils.cell.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(x)[0]))):
        x, y = a[co].value, b[co].value
        if x != y:
            n += 1
            print(f'{n:3d} [{s}] {co}\n     TU : {x!r}  (cache {va[s][co].value!r})\n     OFI: {y!r}  (cache {vb[s][co].value!r})')
for s in wb.sheetnames:
    if s not in wa.sheetnames: print('SOLO EN OFICIAL:', s, wb[s].sheet_state, wb[s].max_row, wb[s].max_column)
print('TOTAL', n)
